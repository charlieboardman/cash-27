"""Bookkeeper outputs use active expenses and final engine valuation."""
import csv
import io
import re
import uuid

import pytest
from openpyxl import load_workbook

from cash.db import connect
from cash.web.app import create_app


HEADERS = ['date', 'Account', 'mxn', 'ER', 'usd', 'calc_usd',
           'voucher_id', 'receipt_id', 'addendum_id', 'class_code',
           'project', 'recipient', 'vendor', 'description', 'notes']


CASH_HEADERS = ['date', 'Activity', 'From', 'To', 'mxn', 'ER', 'usd', 'calc_usd',
                'Safe Balance USD', 'Safe Balance MXN', *HEADERS[6:]]

@pytest.fixture
def env(tmp_path):
    app = create_app(tmp_path, admin_password='test', testing=True)
    client = app.test_client()
    token = client.post('/api/login', json={'password': 'test'}).json['data']['csrf']
    with connect(tmp_path / 'cash.db') as conn:
        yield client, {'X-CSRF': token}, conn


def insert(conn, **fields):
    values = dict(movement_id=str(uuid.uuid4()), submission_id=str(uuid.uuid4()),
                  date='2026-09-02', same_day_seq=0, txn_type='expense',
                  from_account='SAFE', to_account='EXP_SUPPLIES',
                  usd_cents=0, mxn_centavos=0)
    values.update(fields)
    return conn.execute(
        f"INSERT INTO transactions ({','.join(values)}) VALUES ({','.join('?' for _ in values)})",
        list(values.values())).lastrowid


def opening(conn, *, mxn=155000, rate=15500):
    return insert(conn, date='2026-08-01', txn_type='opening',
                  from_account='OPENING_EQUITY', to_account='SAFE',
                  usd_cents=10000, mxn_centavos=mxn,
                  assigned_rate_milli=rate if mxn else None)


def preview(env, year=2027, export_type="expenses"):
    response = env[0].get(f'/api/export/transactions.csv?fy={year}&export_type={export_type}')
    assert response.status_code == 200
    reader = csv.DictReader(io.StringIO(response.data.decode()))
    assert reader.fieldnames == (HEADERS if export_type == "expenses" else CASH_HEADERS)
    return list(reader)


def workbook(env, query=''):
    response = env[0].get('/api/export/transactions.xlsx' + query)
    assert response.status_code == 200
    assert re.search(r'all-years-expenses-\d{8}T\d{6}Z\.xlsx', response.headers['Content-Disposition'])
    return load_workbook(io.BytesIO(response.data))


def test_csv_and_xlsx_include_only_active_expenses(env):
    client, headers, conn = env
    opening(conn)
    conn.execute("INSERT INTO accounts (account_id, name, kind) VALUES ('SIGNOUT', 'Person', 'signout')")
    for kind, source, target, usd, mxn in [
        ('bank_withdrawal', 'BANK', 'SAFE', 1000, 0),
        ('donation', 'DONATION', 'SAFE', 1000, 0),
        ('signout', 'SAFE', 'SIGNOUT', 500, 0),
        ('addendum', 'SAFE', 'SIGNOUT', 100, 0),
        ('exchange', 'SAFE', 'FX_COUNTERPARTY', 1000, -15500),
        ('adjustment', 'SAFE', 'RECONCILIATION', 100, 0),
    ]:
        insert(conn, txn_type=kind, from_account=source, to_account=target,
               usd_cents=usd, mxn_centavos=mxn,
               same_day_seq=conn.execute('SELECT count(*) FROM transactions').fetchone()[0])
    insert(conn, date='2026-09-03', usd_cents=9900, posting_state='superseded')
    insert(conn, date='2026-09-04', usd_cents=8800, posting_state='void')
    # Insert out of date order to check report ordering independently of IDs.
    insert(conn, date='2027-09-01', usd_cents=200, description='Next FY')
    insert(conn, date='2026-09-06', usd_cents=-100, description='Refund')
    insert(conn, date='2026-09-05', mxn_centavos=10000, description='Expense')
    rows = preview(env)
    assert [r['description'] for r in rows] == ['Expense', 'Refund']
    assert [r['calc_usd'] for r in rows] == ['6.45', '-1.00']
    assert rows[0]['ER'] == '15.500'  # Not 100 / the rounded 6.45.
    assert rows[1]['ER'] == ''
    cash_rows = preview(env, export_type='cash_ledger')
    assert [r['Activity'] for r in cash_rows] == [
        'Bank withdrawal', 'Donation', 'Cash signed out', 'Additional cash signed out',
        'Currency exchange', 'Adjustment', 'Expense', 'Expense']
    assert [(r['Safe Balance USD'], r['Safe Balance MXN']) for r in cash_rows[-2:]] == [
        ('103.00', '1605.00'), ('104.00', '1605.00')]
    wb = workbook(env)
    assert wb.sheetnames == ['expenses', 'FY2027', 'FY2028', 'Export']
    assert list(wb['expenses'].values)[0] == tuple(HEADERS)
    assert [r[HEADERS.index('description')] for r in list(wb['expenses'].values)[1:]] == ['Expense', 'Refund', 'Next FY']
    assert wb['FY2027']['C2'].value == 100
    assert wb['FY2027']['D2'].value == 15.5
    assert wb['FY2027']['F2'].value == 6.45
    assert wb['FY2027']['F3'].value == -1
    assert wb['FY2027']['D3'].value is None
    for column in ('C', 'E', 'F'):
        assert wb['FY2027'][column + '2'].data_type == 'n'
        assert wb['FY2027'][column + '2'].number_format == '0.00'
    assert wb['FY2027']['D2'].number_format == '0.000'


@pytest.mark.parametrize('usd,mxn,face,expected,balances', [
    (500, 23250, 31000, ('15.500', '20.00'), ('95.00', '1317.50')),
    (2000, -7000, 24000, ('15.500', '15.48'), ('80.00', '1620.00')),
    (125, 0, None, ('', '1.25'), ('98.75', '1550.00')),
])
def test_rate_uses_peso_valuation_for_mixed_tender_and_change(env, usd, mxn, face, expected, balances):
    opening(env[2])
    insert(env[2], usd_cents=usd, mxn_centavos=mxn, receipt_face_mxn_centavos=face)
    row = preview(env)[0]
    assert (row['ER'], row['calc_usd']) == expected
    cash_row = preview(env, export_type='cash_ledger')[0]
    assert (cash_row['Safe Balance USD'], cash_row['Safe Balance MXN']) == balances


def test_rounds_half_cents_only_at_export_and_preserves_code_text(env):
    conn = env[2]
    opening(conn, rate=8000)
    conn.execute("INSERT INTO accounts (account_id, name, kind) VALUES ('0010', 'Supplies', 'expense')")
    insert(conn, to_account='0010', mxn_centavos=100, vendor='=1+1')
    row = preview(env)[0]
    assert row['calc_usd'] == '0.13'
    assert row['ER'] == '8.000'
    assert row['vendor'] == "'=1+1"
    wb = workbook(env)
    assert wb['expenses']['B2'].value == '0010'
    assert wb['expenses']['B2'].data_type == 's'
    assert wb['expenses']['M2'].value == "'=1+1"
    assert wb['expenses']['M2'].data_type == 's'


def test_unresolved_rate_stays_blank_then_uses_full_later_funding(env):
    client, headers, conn = env
    opening(conn)
    tid = insert(conn, mxn_centavos=200000)
    before = preview(env)[0]
    assert before['ER'] == before['calc_usd'] == ''
    assert preview(env, export_type='cash_ledger')[0]['Safe Balance MXN'] == '-450.00'
    warned = client.get('/api/export/transactions.csv?fy=2027')
    assert warned.headers['X-Report-Unresolved-Basis'] == '1'
    insert(conn, date='2026-09-03', txn_type='donation', from_account='DONATION',
           to_account='SAFE', mxn_centavos=45000, assigned_rate_milli=15000)
    after = preview(env)[0]
    assert after['calc_usd'] == '130.00'
    assert after['ER'] == '15.385'
    assert preview(env, export_type='cash_ledger')[0]['Safe Balance MXN'] == '-450.00'
    # Filtering the export never filters the history used by the engine.
    wb = workbook(env, f'?since_txn_id={tid - 1}')
    assert wb['expenses']['D2'].value == 15.385
    assert wb['expenses']['F2'].value == 130
    assert len(list(wb['expenses'].values)) == 2


def test_ledger_csv_and_period_xlsx_have_matching_expense_data(env):
    client, headers, conn = env
    opening(conn)
    insert(conn, mxn_centavos=10000, description='This FY')
    insert(conn, date='2027-09-01', usd_cents=200, description='Next FY')
    csv_response = client.get('/api/export/transactions.csv')
    assert csv_response.status_code == 200
    assert re.search(r'all-years-expenses-\d{8}T\d{6}Z\.csv', csv_response.headers['Content-Disposition'])
    rows = list(csv.DictReader(io.StringIO(csv_response.data.decode())))
    assert [r['description'] for r in rows] == ['This FY', 'Next FY']
    assert rows[0]['calc_usd'] == '6.45'
    response = client.get('/api/export/transactions.xlsx?fy=2027')
    assert response.status_code == 200
    assert re.search(r'FY2027-expenses-\d{8}T\d{6}Z\.xlsx', response.headers['Content-Disposition'])
    wb = load_workbook(io.BytesIO(response.data))
    assert wb.sheetnames == ['FY2027', 'Export']
    assert list(wb.active.values)[0] == tuple(HEADERS)
    assert wb.active['N2'].value == 'This FY'
    assert wb.active['D2'].value == 15.5
    assert wb.active['D2'].number_format == '0.000'
    assert wb.active['F2'].value == 6.45
    assert len(list(wb.active.values)) == 2
    info = dict(wb['Export'].values)
    assert info['Export'] == 'Expense report'
    assert info['Fiscal year'] == 'FY2027'
    assert info['Through transaction'] == 3
    assert info['Ledger revision'] == 0
    assert re.fullmatch(r'\d{4}-\d\d-\d\d \d\d:\d\d:\d\d', info['Generated (UTC)'])
    assert 'After transaction' not in info
    since = load_workbook(io.BytesIO(client.get('/api/export/transactions.xlsx?since_txn_id=2').data))
    assert dict(since['Export'].values)['After transaction'] == 2
    assert dict(since['Export'].values)['Fiscal year'] == 'All years'


@pytest.mark.parametrize('format', ['csv', 'xlsx'])
def test_safe_balances_follow_settlements_and_full_history_despite_cursor(env, format):
    client, _, conn = env
    opening(conn)
    conn.execute("INSERT INTO accounts (account_id, name, kind) VALUES ('SIGNOUT', 'Person', 'signout')")
    insert(conn, date='2026-09-01', txn_type='signout', to_account='SIGNOUT',
           usd_cents=3000, mxn_centavos=10000)
    # Insert the return before the earlier expense to exercise replay chronology.
    insert(conn, same_day_seq=1, txn_type='addendum', from_account='SIGNOUT',
           to_account='SAFE', usd_cents=500, mxn_centavos=2000)
    cutoff = insert(conn, from_account='SIGNOUT', usd_cents=1000, mxn_centavos=3000,
                    description='Signee expense')
    insert(conn, same_day_seq=2, usd_cents=125, mxn_centavos=200,
           description='Direct expense')
    movement = str(uuid.uuid4())
    for line, amount in [(2, 200), (1, 100)]:
        insert(conn, same_day_seq=3, movement_id=movement, allocation_line=line,
               usd_cents=amount, mxn_centavos=amount, description=f'Split {line}')
    insert(conn, date='2026-09-03', txn_type='bank_withdrawal',
           from_account='BANK', to_account='SAFE', usd_cents=1000)
    rows = [r for r in preview(env, export_type='cash_ledger') if r['Activity'] == 'Expense']
    assert [r['description'] for r in rows] == [
        'Signee expense', 'Direct expense', 'Split 1', 'Split 2']
    assert [(r['Safe Balance USD'], r['Safe Balance MXN']) for r in rows] == [
        ('70.00', '1450.00'), ('73.75', '1468.00'),
        ('70.75', '1465.00'), ('70.75', '1465.00')]
    response = client.get(f'/api/export/transactions.{format}?since_txn_id={cutoff}&export_type=cash_ledger')
    assert response.status_code == 200
    if format == 'csv':
        values = list(csv.DictReader(io.StringIO(response.data.decode())))
        assert values[:-1] == rows[1:]
        assert values[-1]['Activity'] == 'Bank withdrawal'
        assert values[-1]['Safe Balance USD'] == '80.75'
    else:
        wb = load_workbook(io.BytesIO(response.data))
        values = [dict(zip(CASH_HEADERS, row)) for row in list(wb['cash-ledger'].values)[1:]]
        assert [r['description'] for r in values[:-1]] == [r['description'] for r in rows[1:]]
        assert [(r['Safe Balance USD'], r['Safe Balance MXN']) for r in values] == [
            (73.75, 1468), (70.75, 1465), (70.75, 1465), (80.75, 1465)]


@pytest.mark.parametrize('format', ['csv', 'xlsx'])
def test_cash_ledger_includes_latest_withdrawal_and_exchange(env, format):
    client, headers, conn = env
    opening(conn)
    insert(conn, usd_cents=500)
    withdrawal = insert(conn, date='2026-09-03', txn_type='bank_withdrawal',
                        from_account='BANK', to_account='SAFE', usd_cents=20000)
    insert(conn, date='2026-09-04', txn_type='exchange', to_account='FX_COUNTERPARTY',
           usd_cents=1000, mxn_centavos=-15500)
    insert(conn, date='2026-09-05', usd_cents=9900, posting_state='superseded')
    insert(conn, date='2026-09-06', usd_cents=8800, posting_state='void')
    assert len(preview(env)) == 1
    response = client.get(f'/api/export/transactions.{format}?export_type=cash_ledger')
    assert response.status_code == 200
    assert re.search(rf'all-years-cash-ledger-\d{{8}}T\d{{6}}Z\.{format}', response.headers['Content-Disposition'])
    if format == 'csv':
        rows = list(csv.DictReader(io.StringIO(response.data.decode())))
        assert rows[-2]['Safe Balance USD'] == '295.00'
        assert rows[-1]['Safe Balance USD'] == '285.00'
        assert rows[-1]['Safe Balance MXN'] == '1705.00'
        assert rows[-1]['ER'] == '15.500'
        assert rows[-1]['mxn'] == '-155.00'
        assert rows[-2]['calc_usd'] == rows[-1]['calc_usd'] == ''
    else:
        wb = load_workbook(io.BytesIO(response.data))
        assert wb.sheetnames == ['cash-ledger', 'FY2026', 'FY2027', 'Export']
        rows = [dict(zip(CASH_HEADERS, r)) for r in list(wb['cash-ledger'].values)[1:]]
        assert rows[-2]['Safe Balance USD'] == 295
        assert rows[-1]['Safe Balance USD'] == 285
        assert rows[-1]['Safe Balance MXN'] == 1705
        assert rows[-1]['ER'] == 15.5
        assert rows[-1]['mxn'] == -155
        assert rows[-2]['calc_usd'] is rows[-1]['calc_usd'] is None
        assert wb['cash-ledger']['I4'].data_type == 'n'
        assert wb['cash-ledger']['I4'].number_format == '0.00'
    assert [r['Activity'] for r in rows] == [
        'Opening balance', 'Expense', 'Bank withdrawal', 'Currency exchange']
    assert (rows[-2]['From'], rows[-2]['To']) == ('BANK', 'SAFE')
    # The period excludes the opening row, but keeps its cash in the balances.
    report = preview(env, export_type='cash_ledger')
    assert len(report) == 3
    assert report[-1]['Safe Balance USD'] == '285.00'


@pytest.mark.parametrize('export_type', ['invalid', ''])
def test_invalid_export_type_is_rejected(env, export_type):
    response = env[0].get(f'/api/export/transactions.csv?fy=2027&export_type={export_type}')
    assert response.status_code == 400
    assert response.json['error']['code'] == 'bad_export_type'
