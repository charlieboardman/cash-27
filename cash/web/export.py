"""Expense reports and cash ledgers from the full authoritative replay."""

from __future__ import annotations

import csv
from datetime import datetime, timezone
import io
from decimal import Decimal, ROUND_HALF_UP

from flask import Blueprint, g, request, send_file
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell

from .. import repo
from ..prep import fiscal_year_for
from .app import admin_required, err

bp = Blueprint("export", __name__)

COLUMNS = ["date", "Account", "mxn", "ER", "usd", "calc_usd",
           "voucher_id", "receipt_id", "addendum_id", "class_code",
           "project", "recipient", "vendor", "description", "notes"]

CASH_COLUMNS = ["date", "Activity", "From", "To", "mxn", "ER", "usd", "calc_usd",
                "Safe Balance USD", "Safe Balance MXN", *COLUMNS[6:]]
EXPORT_TYPES = {"expenses": COLUMNS, "cash_ledger": CASH_COLUMNS}
EXPORT_NAMES = {"expenses": "expenses", "cash_ledger": "cash-ledger"}
EXPORT_LABELS = {"expenses": "Expense report", "cash_ledger": "Cash ledger"}

PREFIXES = ("=", "+", "-", "@")


def _escape(v):
    if isinstance(v, str) and v[:1] in PREFIXES:
        return "'" + v
    return v


def export_rows(conn, export_type="expenses", since=None):
    """Select presentation rows; valuation and balances always replay all history."""
    sql = "SELECT t.* FROM transactions t WHERE t.posting_state='active'"
    if export_type == "expenses":
        sql += " AND t.to_account IN (SELECT account_id FROM accounts WHERE kind='expense')"
    params = []
    if since is not None:
        sql += " AND t.txn_id > ?"
        params.append(since)
    sql += " ORDER BY t.date, t.same_day_seq, t.movement_id, t.allocation_line, t.txn_id"
    return conn.execute(sql, params).fetchall()


def _activity(t):
    if t["txn_type"] == "addendum":
        return "Cash returned" if t["to_account"] == "SAFE" else "Additional cash signed out"
    return {"expense": "Expense", "bank_withdrawal": "Bank withdrawal",
            "exchange": "Currency exchange", "opening": "Opening balance",
            "donation": "Donation", "signout": "Cash signed out",
            "adjustment": "Adjustment"}.get(t["txn_type"], t["txn_type"].replace("_", " ").title())


def _rounded(value, places):
    result = value.quantize(Decimal(places), rounding=ROUND_HALF_UP)
    return abs(result) if result == 0 else result


def _row_values(t, res):
    mid = t["movement_id"]
    mres = res.movements.get(mid)
    safe = mres.safe_balance_after if mres else None
    calc = None
    rate = None
    usd = Decimal(t["usd_cents"]) / 100
    mxn = Decimal(t["mxn_centavos"]) / 100
    if mres:
        calc = mres.calc_usd_by_row.get(str(t["txn_id"]))
        if str(t["txn_id"]) in mres.unresolved_rows:
            calc = None
        # Final replay valuation includes basis supplied by later funding.
        # The initial basis_out_by_row alone omits those resolved claims.
        if calc is not None and mxn and (basis := calc - usd):
            implied = mxn / basis
            if implied > 0:
                rate = _rounded(implied, "0.001")
    if t["txn_type"] == "exchange" and usd > 0 and mxn < 0:
        rate = _rounded(-mxn / usd, "0.001")
    return {
        **dict(t), "Account": t["to_account"], "Activity": _activity(t),
        "From": t["from_account"], "To": t["to_account"],
        "mxn": _rounded(mxn, "0.01"), "ER": rate,
        "usd": _rounded(usd, "0.01"),
        "calc_usd": None if calc is None else _rounded(calc, "0.01"),
        "Safe Balance USD": None if safe is None else _rounded(Decimal(safe.usd) / 100, "0.01"),
        "Safe Balance MXN": None if safe is None else _rounded(Decimal(safe.mxn) / 100, "0.01"),
    }


MIMETYPES = {
    "csv": "text/csv",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}
NUMBER_FORMATS = {"mxn": "0.00", "ER": "0.000", "usd": "0.00", "calc_usd": "0.00",
                  "Safe Balance USD": "0.00", "Safe Balance MXN": "0.00"}


def export_values(txns, replay_result, export_type="expenses"):
    rows = []
    for txn in txns:
        values = _row_values(txn, replay_result)
        rows.append([_escape(values[column]) for column in EXPORT_TYPES[export_type]])
    return rows


def render_csv(rows, *, columns=COLUMNS):
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(columns)
    writer.writerows(rows)
    return buf.getvalue().encode()


def render_xlsx(rows, *, columns=COLUMNS, sheet_name="expenses", fy_month=None,
                info=()):
    """info rows (label, value) go on a closing "Export" sheet."""
    wb = Workbook(write_only=True)
    all_ws = wb.create_sheet(sheet_name)
    all_ws.append(columns)
    by_fy = {}

    def append_row(ws, values):
        cells = []
        for column, value in zip(columns, values):
            cell = WriteOnlyCell(ws, value=value)
            if column in NUMBER_FORMATS:
                cell.number_format = NUMBER_FORMATS[column]
            cells.append(cell)
        ws.append(cells)

    for row in rows:
        append_row(all_ws, row)
        if fy_month is not None:
            fy = fiscal_year_for(row[columns.index("date")], fy_month)
            by_fy.setdefault(fy, []).append(row)
    for fy, values in sorted(by_fy.items()):
        ws = wb.create_sheet(f"FY{fy}")
        ws.append(columns)
        for row in values:
            append_row(ws, row)
    if info:
        ws = wb.create_sheet("Export")
        for label, value in info:
            ws.append([label, value])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def warnings(rows, replay_result) -> tuple[int, int]:
    """Unallocated expenses and unresolved MXN basis among the exported rows."""
    unallocated = sum(1 for txn in rows if txn["to_account"] == "UNALLOCATED_EXPENSE")
    movement_ids = {txn["movement_id"] for txn in rows}
    unresolved = sum(1 for claim in replay_result.open_claims
                     if claim.movement_id in movement_ids)
    return unallocated, unresolved


@bp.get("/api/export/transactions.csv", defaults={"format": "csv"})
@bp.get("/api/export/transactions.xlsx", defaults={"format": "xlsx"})
@admin_required
def export_transactions(format):
    """One fiscal year (fy=2027) or every year (no fy). since_txn_id limits
    the rows, never the history used for valuation and balances."""
    from flask import current_app
    fy_month = current_app.config["FY_START_MONTH"]
    export_type = request.args.get("export_type", "expenses")
    if export_type not in EXPORT_TYPES:
        return err("bad_export_type", "Choose Expense report or Cash ledger.", 400)
    fy = request.args.get("fy", "")
    if fy and (not fy.isdigit() or not 2000 <= int(fy) <= 2200):
        return err("bad_fy", "Choose a fiscal year between 2000 and 2200.", 400)
    since = request.args.get("since_txn_id", type=int)
    g.db.execute("BEGIN")
    try:
        res = repo.replay_projected(g.db)
        txns = [t for t in export_rows(g.db, export_type, since)
                if not fy or fiscal_year_for(t["date"], fy_month) == int(fy)]
        through = g.db.execute("SELECT COALESCE(MAX(txn_id), 0) FROM transactions").fetchone()[0]
        revision = repo.ledger_revision(g.db)
    finally:
        g.db.rollback()
    rows = export_values(txns, res, export_type)
    unallocated, unresolved = warnings(txns, res)
    columns = EXPORT_TYPES[export_type]
    generated = datetime.now(timezone.utc)
    period = f"FY{fy}" if fy else "all-years"
    name = f"{period}-{EXPORT_NAMES[export_type]}-{generated:%Y%m%dT%H%M%SZ}"
    if format == "csv":
        content = render_csv(rows, columns=columns)
    else:
        info = [("Export", EXPORT_LABELS[export_type]),
                ("Fiscal year", f"FY{fy}" if fy else "All years"),
                ("Generated (UTC)", generated.strftime("%Y-%m-%d %H:%M:%S")),
                ("Through transaction", through),
                ("Ledger revision", revision)]
        if since is not None:
            info.append(("After transaction", since))
        content = render_xlsx(rows, columns=columns, info=info,
                              sheet_name=period if fy else EXPORT_NAMES[export_type],
                              fy_month=None if fy else fy_month)
    response = send_file(io.BytesIO(content), mimetype=MIMETYPES[format],
                         as_attachment=True, download_name=f"{name}.{format}")
    response.headers["X-Through-Txn-Id"] = str(through)
    response.headers["X-Ledger-Revision"] = str(revision)
    response.headers["X-Report-Unallocated"] = str(unallocated)
    response.headers["X-Report-Unresolved-Basis"] = str(unresolved)
    return response
