"""Read-only bookkeeper views and portable document packages. No retained state."""
from collections import defaultdict
from datetime import date, datetime, timezone
from decimal import Decimal
import hashlib
import io
import json
from pathlib import Path
import re
import tempfile
from urllib.parse import quote
from zipfile import ZipFile, ZIP_STORED

from flask import Blueprint, current_app, g, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.hyperlink import Hyperlink

from .. import repo
from ..prep import fiscal_year_for
from .app import admin_required, err
from .export import export_rows, _row_values, _escape, render_csv, EXPORT_TYPES

bp = Blueprint('bookkeeper', __name__)
BASE_COLUMNS = ['Date', 'Payee/Vendor', 'Category', 'Class', 'Description',
                'MXN', 'ER', 'USD', 'Calc USD']
MONEY = {'MXN', 'USD', 'Calc USD', 'Safe Balance USD', 'Safe Balance MXN'}


def options():
    kind = request.args.get('export_type', 'expenses')
    if kind not in EXPORT_TYPES:
        raise ValueError('Choose Expense report or Cash ledger.')
    year = request.args.get('fy', '')
    if year and (not year.isdigit() or not 2000 <= int(year) <= 2200):
        raise ValueError('Enter a fiscal year between 2000 and 2200.')
    start, end = request.args.get('start', ''), request.args.get('end', '')
    for value in (start, end):
        if value:
            if date.fromisoformat(value).isoformat() != value:
                raise ValueError('Use YYYY-MM-DD dates.')
    if start and end and start > end:
        raise ValueError('The end date must follow the start date.')
    return dict(export_type=kind, fy=year, start=start, end=end)


class DocumentNameClash(Exception):
    """Two different documents would get the same file name in a package."""


def _path(root, relative):
    path = (root / relative).resolve()
    if path.parent != (root / 'documents').resolve() or path.suffix.lower() != '.pdf':
        return None
    return path


def build_report(conn, root, fy_month, opts):
    """All values and document associations come from the caller's read snapshot."""
    txns = [dict(t) for t in export_rows(conn, opts['export_type'])
            if (not opts['fy'] or fiscal_year_for(t['date'], fy_month) == int(opts['fy']))
            and (not opts['start'] or t['date'] >= opts['start'])
            and (not opts['end'] or t['date'] <= opts['end'])]
    replayed = repo.replay_projected(conn)
    accounts = {r['account_id']: dict(r) for r in conn.execute('SELECT * FROM accounts')}
    classes = {r['class_code']: r['name'] for r in conn.execute('SELECT * FROM classes')}
    vouchers = {r['voucher_id']: dict(r) for r in conn.execute('SELECT * FROM vouchers')}
    receipts = {r['receipt_id']: dict(r) for r in conn.execute("SELECT * FROM receipts WHERE status!='void'")}
    addenda = {r['addendum_id']: dict(r) for r in conn.execute("SELECT * FROM addenda WHERE status!='void'")}
    documents = {r['document_id']: dict(r) for r in conn.execute('SELECT * FROM documents')}
    numbers = repo.receipt_numbers(conn)
    direct, family = defaultdict(list), defaultdict(list)
    for row in conn.execute(
            "SELECT td.*, t.voucher_id, t.addendum_id, t.receipt_id, t.txn_type "
            "FROM transaction_documents td JOIN transactions t USING(txn_id) "
            "WHERE t.posting_state='active' ORDER BY td.position, td.role DESC, td.document_id"):
        ref = repo.document_reference(row, row['role'], row['document_id'], numbers)
        role = 'Addendum' if row['addendum_id'] and row['role'] == 'voucher' else 'Voucher' if row['role'] == 'voucher' else 'Receipt'
        item = (row['document_id'], ref, role)
        direct[row['txn_id']].append(item)
        if row['voucher_id'] and row['txn_type'] in ('signout', 'addendum'):
            family[row['voucher_id']].append(item)
    for aid, a in addenda.items():
        if a['document_id']:
            family[a['voucher_id']].append((a['document_id'], aid, 'Addendum'))

    def named(code, names):
        name = names.get(code, '')
        return f'{code} — {name}' if name and name != code else code or ''

    groups, files, names = {}, {}, {}
    category_names = {code: row['name'] for code, row in accounts.items()}
    columns = BASE_COLUMNS.copy()
    if opts['export_type'] == 'cash_ledger':
        columns[1:1] = ['Activity', 'From', 'To']
        columns.extend(['Safe Balance USD', 'Safe Balance MXN'])
    columns.append('Documents')
    for t in txns:
        raw = _row_values(t, replayed)
        v, r = vouchers.get(t['voucher_id'], {}), receipts.get(t['receipt_id'], {})
        vendor = t['vendor'] or r.get('vendor') or (v.get('recipient') if v.get('type') == 'direct_expense' else '') or ''
        category = t['to_account'] if accounts.get(t['to_account'], {}).get('kind') == 'expense' else ''
        values = {'Date': t['date'], 'Payee/Vendor': vendor,
                  'Category': named(category, category_names), 'Class': named(t['class_code'], classes),
                  'Description': '\n'.join(filter(None, [t['description'] or v.get('purpose') or '', t['notes']])),
                  'MXN': raw['mxn'], 'ER': raw['ER'], 'USD': raw['usd'], 'Calc USD': raw['calc_usd'],
                  'Activity': raw['Activity'], 'From': named(raw['From'], category_names),
                  'To': named(raw['To'], category_names),
                  'Safe Balance USD': raw['Safe Balance USD'], 'Safe Balance MXN': raw['Safe Balance MXN']}
        key = (t['movement_id'], t['receipt_id'])
        group = groups.setdefault(key, {'rows': [], 'documents': {}, 'reference': t['receipt_id'] or t['addendum_id'] or t['voucher_id'] or f"Txn-{t['txn_id']}", 'expected': set()})
        group['rows'].append({'values': values, 'txn_id': t['txn_id'], 'category': category,
                              'class_code': t['class_code'] or '',
                              'context': ' · '.join(filter(None, [t['recipient'] or (v.get('recipient') if v.get('type') == 'signout' else ''), t['project']]))})
        items = list(direct[t['txn_id']]) + family[t['voucher_id']]
        if r.get('document_id'):
            items.append((r['document_id'], t['receipt_id'], 'Receipt'))
        for reference in (t['voucher_id'], t['receipt_id'], t['addendum_id']):
            if reference:
                group['expected'].add(reference)
        for did, reference, role in items:
            if did in group['documents']:
                continue
            d = documents.get(did)
            if not d or d['state'] == 'discarded':
                continue
            path = _path(root, d['rel_path'])
            available = bool(path and path.is_file() and d['state'] != 'missing')
            # Prefer the voucher role when an older upload has both roles.
            doc = files.get(did)
            if doc is None:
                slug = re.sub(r'[^A-Za-z0-9_-]', '-', reference).strip('-') or 'document'
                filename = f"documents/{slug}.pdf"
                if names.setdefault(filename, did) != did:
                    raise DocumentNameClash(
                        f"Documents {names[filename]} and {did} would both be named "
                        f"{slug}.pdf. Correct how they are filed before exporting.")
                doc = {'id': did, 'reference': reference, 'role': role, 'filename': filename,
                       'path': path, 'sha256': d['sha256'], 'available': available,
                       'url': f'/api/documents/{quote(did, safe="")}/stream'}
                files[did] = doc
            group['documents'][did] = doc

    result = list(groups.values())
    for group in result:
        group['documents'] = list(group['documents'].values())
        # Existing identifiers remain visible even when paperwork was never linked.
        linked = {d['reference'] for d in group['documents']}
        group['missing_references'] = sorted(group.pop('expected') - linked)
        group['categories'] = sorted({r['category'] for r in group['rows'] if r['category']})
        group['classes'] = sorted({r['class_code'] for r in group['rows'] if r['class_code']})
        group['search'] = ' '.join(str(v or '') for row in group['rows'] for v in row['values'].values()) + ' ' + group['reference'] + ' ' + ' '.join(r['context'] for r in group['rows']) + ' ' + ' '.join(d['reference'] for d in group['documents'])
        group['sort'] = {c: str(sum((r['values'][c] or Decimal(0) for r in group['rows']), Decimal(0))) if c in MONEY else str(group['rows'][0]['values'].get(c) or '') for c in columns if c != 'Documents'}
    return {'groups': result, 'files': list(files.values()), 'columns': columns, 'options': opts,
            'title': 'Expense report' if opts['export_type'] == 'expenses' else 'Cash ledger',
            'period': f"FY{opts['fy']}" if opts['fy'] else 'All fiscal years',
            'created': datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC'),
            'revision': conn.execute('SELECT ledger_revision FROM ledger_state').fetchone()[0],
            'categories': [(c, named(c, category_names)) for c in sorted({r['category'] for g in result for r in g['rows'] if r['category']})],
            'classes': [(c, named(c, classes)) for c in sorted({r['class_code'] for g in result for r in g['rows'] if r['class_code']})],
            'row_count': len(txns), 'money': MONEY}


def render_report(report, offline=False):
    static = Path(current_app.static_folder)
    return render_template('bookkeeper.html', report=report, offline=offline,
                           report_css=(static / 'css/bookkeeper.css').read_text(),
                           report_js=(static / 'js/bookkeeper.js').read_text())


def _document_text(group):
    return '; '.join([f"{d['reference']}: {d['filename']}" if d['available'] else f"{d['reference']}: Missing PDF" for d in group['documents']] + [f'{ref}: Not attached' for ref in group['missing_references']]) or 'No documents linked'


def spreadsheets(report):
    columns = report['columns']
    wb = Workbook()
    sheet = wb.active
    sheet.title = 'Report'
    sheet.append(columns)
    sheet.freeze_panes = 'F2'
    doc_sheet = wb.create_sheet('Documents')
    doc_sheet.append(['Reference', 'Type', 'PDF', 'Status'])
    doc_rows = {}
    for d in report['files']:
        doc_rows[d['id']] = doc_sheet.max_row + 1
        doc_sheet.append([_escape(d['reference']), d['role'], d['filename'] if d['available'] else '', 'Included' if d['available'] else 'Missing PDF'])
        if d['available']:
            doc_sheet.cell(doc_sheet.max_row, 3).hyperlink = d['filename']
            doc_sheet.cell(doc_sheet.max_row, 3).style = 'Hyperlink'
    csv_rows = []
    for group in report['groups']:
        for row in group['rows']:
            values = [_escape(row['values'].get(c)) if c != 'Documents' else _escape(_document_text(group)) for c in columns]
            csv_rows.append(values)
            sheet.append(values)
            for col, name in enumerate(columns, 1):
                cell = sheet.cell(sheet.max_row, col)
                if name in MONEY or name == 'ER':
                    cell.number_format = '0.000' if name == 'ER' else '0.00'
                if name == 'Calc USD':
                    cell.fill = PatternFill('solid', fgColor='EEF5EF')
                    cell.font = Font(bold=True)
            if group['documents']:
                cell = sheet.cell(sheet.max_row, len(columns))
                first = doc_rows[group['documents'][0]['id']]
                cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'Documents'!A{first}")
                cell.style = 'Hyperlink'
    for ws in (sheet, doc_sheet):
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='293C32')
        for column in ws.columns:
            ws.column_dimensions[column[0].column_letter].width = min(48, max(14, len(str(column[0].value)) + 3))
    buf = io.BytesIO()
    wb.save(buf)
    return render_csv(csv_rows, columns=columns), buf.getvalue()


def _load():
    opts = options()
    g.db.execute('BEGIN')
    try:
        return build_report(g.db, current_app.config['DATA_ROOT'], current_app.config['FY_START_MONTH'], opts)
    finally:
        g.db.rollback()


@bp.get('/reports/bookkeeper')
@admin_required
def page():
    try:
        report = _load()
    except ValueError as exc:
        return err('bad_report_filter', str(exc), 400)
    except DocumentNameClash as exc:
        return err('document_name_clash', str(exc), 409)
    return render_report(report)


@bp.get('/api/export/bookkeeper.zip')
@admin_required
def package():
    try:
        report = _load()
    except ValueError as exc:
        return err('bad_report_filter', str(exc), 400)
    except DocumentNameClash as exc:
        return err('document_name_clash', str(exc), 409)
    # PDFs are already compressed. Build on disk, with bounded copy buffers.
    output = tempfile.TemporaryFile(dir=current_app.config['DATA_ROOT'] / 'staging')
    try:
        csv_bytes, xlsx_bytes = spreadsheets(report)
        with ZipFile(output, 'w', compression=ZIP_STORED) as archive:
            archive.writestr('report.csv', csv_bytes)
            archive.writestr('report.xlsx', xlsx_bytes)
            archive.writestr('index.html', render_report(report, offline=True))
            archive.writestr('README.txt', 'Extract the entire ZIP into one folder, then open index.html or report.xlsx.\nPDF links use the documents folder. No login or internet is needed.\nThis is a working snapshot, not a retained final report or a QBO import file.\nMissing or unattached documents are marked in the report.\nSearch/category/class filters affect the browser view only; this package contains the selected period.\n')
            manifest = {'created': report['created'], 'ledger_revision': report['revision'], 'filters': report['options'], 'documents': [], 'rows': []}
            for d in report['files']:
                if d['available']:
                    digest = hashlib.sha256()
                    with d['path'].open('rb') as source, archive.open(d['filename'], 'w', force_zip64=True) as target:
                        while block := source.read(1024 * 1024):
                            digest.update(block)
                            target.write(block)
                    if digest.hexdigest() != d['sha256']:
                        raise ValueError('A supporting PDF failed its integrity check. No package was delivered.')
                manifest['documents'].append({k: d[k] for k in ('id', 'reference', 'filename', 'sha256', 'available')})
            for group in report['groups']:
                for row in group['rows']:
                    manifest['rows'].append({'txn_id': row['txn_id'], 'documents': [d['id'] for d in group['documents']], 'unattached_references': group['missing_references']})
            archive.writestr('manifest.json', json.dumps(manifest, indent=2))
        output.seek(0)
        response = send_file(output, mimetype='application/zip', as_attachment=True,
                             download_name=f"{report['period'].replace(' ', '-')}-{report['options']['export_type']}-with-documents.zip")
        response.call_on_close(output.close)
        return response
    except (OSError, ValueError) as exc:
        output.close()
        return err('document_package_failed', str(exc), 410)
    except Exception:
        output.close()
        raise
